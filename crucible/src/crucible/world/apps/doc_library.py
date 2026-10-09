"""Document library App: a folder of business documents (policies, contracts, spreadsheets, PDFs).

Tools: list_docs(path=".", pattern=None), read_doc(path, offset=1, limit=400), search_docs(pattern, glob=None),
and write_doc(path, content) when ``params.writable``. Office/PDF documents are rendered to text with
the optional extractors MiMo's verifier uses (python-docx, openpyxl, python-pptx, pdfplumber); when a
library is missing the tool says so instead of returning bytes. State: ``fs`` adapter over ``root``.
"""

from __future__ import annotations

import shlex
import tarfile
import tempfile
from pathlib import Path
from typing import Any

from crucible.core.schema.actions import ToolSpec
from crucible.world.apps.base import App, ToolError, confine
from crucible.world.state.base import snippet

WS_SNIPPET = snippet("ws")
FS_SNIPPET = snippet("fs")

_EXTRACT = r'''
import os
def _docx(p):
    import docx
    d = docx.Document(p)
    out = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)
def _xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)
def _pptx(p):
    import pptx
    prs = pptx.Presentation(p)
    out = []
    for i, s in enumerate(prs.slides, 1):
        out.append(f"## slide {i}")
        for sh in s.shapes:
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out)
def _pdf(p):
    try:
        import pdfplumber
        with pdfplumber.open(p) as pdf:
            return "\n".join((pg.extract_text() or "") for pg in pdf.pages)
    except ImportError:
        import fitz
        return "\n".join(pg.get_text() for pg in fitz.open(p))
EX = {".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".pptx": _pptx, ".pdf": _pdf}
def main(a):
    p = a["path"]; ext = os.path.splitext(p)[1].lower()
    if ext in EX:
        try:
            text = EX[ext](p)
        except ImportError as e:
            return {"text": f"<{ext} document; text extractor unavailable: {e.name}>", "extracted": False}
        lines = text.splitlines()
    else:
        with open(p, "rb") as f:
            raw = f.read()
        if b"\0" in raw[:8192]:
            return {"text": f"<binary file, {len(raw)} bytes>", "extracted": False}
        lines = raw.decode("utf-8", "replace").splitlines()
    off, lim = max(1, int(a.get("offset", 1))), int(a.get("limit", 400))
    chunk = lines[off - 1: off - 1 + lim]
    body = "\n".join(f"{i}\t{l[:2000]}" for i, l in enumerate(chunk, start=off))
    if off - 1 + lim < len(lines):
        body += f"\n... ({len(lines)} lines total; continue with offset={off + lim})"
    return {"text": body, "extracted": True}
'''


class DocLibraryApp(App):
    template = "common/doc_library"

    @property
    def root(self) -> str:
        return self.state_params["root"]

    async def provision(self) -> None:
        await self.ctx.exec.run(f"mkdir -p {shlex.quote(self.root)}")
        seed = self.spec.params.get("seed_dir")
        if seed:
            src = self.ctx.host_seed(seed)
            if self.ctx.exec.is_local:
                import shutil

                shutil.copytree(src, self.root, dirs_exist_ok=True, symlinks=True)
                return
            with tempfile.TemporaryDirectory() as d:
                tar = Path(d) / "seed.tar"
                with tarfile.open(tar, "w") as tf:
                    for p in sorted(src.rglob("*")):
                        tf.add(p, arcname=str(p.relative_to(src)), recursive=False)
                remote = f"/tmp/.crucible_seed_{self.id}.tar"
                await self.ctx.exec.push_file(tar, remote)
                await self.ctx.exec.run(f"tar -xf {remote} -C {shlex.quote(self.root)} && rm -f {remote}")

    async def list_tools(self) -> list[ToolSpec]:
        s, i = {"type": "string"}, {"type": "integer"}
        tools = [
            ToolSpec(app=self.id, name="list_docs", read_only=True, description="List documents in the library.", input_schema={"type": "object", "properties": {"path": s, "pattern": s}}),
            ToolSpec(app=self.id, name="read_doc", read_only=True, description="Read a document as text (docx/xlsx/pptx/pdf are converted); use offset/limit for long ones.",
                     input_schema={"type": "object", "properties": {"path": s, "offset": i, "limit": i}, "required": ["path"]}),
            ToolSpec(app=self.id, name="search_docs", read_only=True, description="Regex search across plain-text documents.", input_schema={"type": "object", "properties": {"pattern": s, "glob": s}, "required": ["pattern"]}),
        ]
        if self.spec.params.get("writable"):
            tools.append(ToolSpec(app=self.id, name="write_doc", description="Create or overwrite a text document.", input_schema={"type": "object", "properties": {"path": s, "content": s}, "required": ["path", "content"]}))
        return tools

    async def tool_list_docs(self, path: str = ".", pattern: str | None = None) -> Any:
        r = await self.policy_exec.run_python(FS_SNIPPET, {"op": "list", "root": confine(self.root, path), "git": False, "include": [], "exclude": [], "pattern": pattern})
        return "\n".join(r["paths"]) or "(no documents)"

    async def tool_read_doc(self, path: str, offset: int = 1, limit: int = 400) -> Any:
        try:
            r = await self.policy_exec.run_python(_EXTRACT, {"path": confine(self.root, path), "offset": offset, "limit": min(limit, 2000)})
        except Exception as e:
            raise ToolError(str(e).split("python snippet failed: ")[-1][:500]) from e
        return r["text"]

    async def tool_search_docs(self, pattern: str, glob: str | None = None) -> Any:
        r = await self.policy_exec.run_python(FS_SNIPPET, {"op": "grep", "root": self.root, "git": False, "include": [], "exclude": [], "pattern": pattern, "glob": glob, "max_hits": 200})
        return "\n".join(f"{p}:{n}:{t}" for p, n, t in r["hits"]) or "no matches"

    async def tool_write_doc(self, path: str, content: str) -> Any:
        if not self.spec.params.get("writable"):
            raise ToolError("the library is read-only")
        await self.policy_exec.run_python(WS_SNIPPET, {"op": "write", "path": confine(self.root, path), "content": content})
        return f"saved {path}"
